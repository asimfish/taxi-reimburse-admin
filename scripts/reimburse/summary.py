import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypinyin import lazy_pinyin

from .checks import check_submission, collect_files
from .layout import apply_main_layout, apply_page_setup
from .printpack import build_invoice_pack, build_print_pack, build_taxi_pack
from .report_text import report_to_text

HEAD_FILL = PatternFill("solid", fgColor="FF31465F")
TOTAL_FILL = PatternFill("solid", fgColor="FFE9EEF4")
ARIAL = "Arial"


def month_title(month):
    return f"{month[:4]}年{int(month[5:])}月"


def pinyin_key(name):
    return "".join(lazy_pinyin(name))


def order_records(persons):
    """按乘坐人拼音排序；每人先出租车（按标签），再网约车（按发票分组、行程单序号）。"""
    records = []
    for p in sorted(persons, key=lambda p: pinyin_key(p["name"])):
        rows = p["report"]["rows"]
        taxi = sorted([r for r in rows if r["type"] == "出租车"], key=lambda r: (r.get("label", ""), r["date"] or ""))
        plat = [r for r in rows if r["type"] == "网约车"]
        file_order = []
        for r in plat:
            f = r["evidence"].get("itinerary", {}).get("file", "~")
            if f not in file_order:
                file_order.append(f)
        file_order.sort(key=lambda f: (pinyin_key(next((r.get("provider") or "" for r in plat if r["evidence"].get("itinerary", {}).get("file", "~") == f), "")), f))
        plat.sort(key=lambda r: (file_order.index(r["evidence"].get("itinerary", {}).get("file", "~")), r["evidence"].get("itinerary", {}).get("no", 999), r["date"] or ""))
        others = [r for r in rows if r["type"] not in ("出租车", "网约车")]
        for r in taxi + plat + others:
            records.append({**r, "person": p["name"]})
    for i, r in enumerate(records, 1):
        r["seq"] = i
    return records


def remark_of(r):
    return r.get("label", "") + (f"【{'；'.join(r['notes'])}】" if r.get("notes") else "")


def _style_header(ws, row, ncol):
    for c in ws[row][:ncol]:
        c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_workbook(month, persons, records, path, unit):
    wb = Workbook()
    ws = wb.active
    ws.title = "行程汇总"
    ws["A1"] = f"交通费报销明细表（{month_title(month)}）"
    ws["A1"].font = Font(name=ARIAL, size=14, bold=True)
    ws["A2"] = f"单位：{unit}"
    ws["A2"].font = Font(name=ARIAL, size=10)
    ws.append(["序号", "日期", "起止地", "公务事由", "乘坐人", "金额", "备注"])
    _style_header(ws, 3, 7)
    for r in records:
        d = dt.date.fromisoformat(r["date"]) if r["date"] else None
        ws.append([r["seq"], d, f"{r['start']} | {r['end']}", r["reason"], r["person"], r["amount"], remark_of(r)])
        row = ws.max_row
        for c in ws[row]:
            c.font = Font(name=ARIAL, size=10)
        ws[f"A{row}"].alignment = Alignment(horizontal="center")
        ws[f"B{row}"].alignment = Alignment(horizontal="center")
        ws[f"B{row}"].number_format = "yyyy-mm-dd"
        ws[f"E{row}"].alignment = Alignment(horizontal="center")
        ws[f"F{row}"].alignment = Alignment(horizontal="right")
        ws[f"F{row}"].number_format = "#,##0.00"
    last = ws.max_row
    ws.append(["合计", None, None, None, None, f"=SUM(F4:F{last})" if records else 0, "按逐笔申报金额合计"])
    t = ws.max_row
    for c in ws[t]:
        c.font = Font(name=ARIAL, size=10, bold=True)
        c.fill = TOTAL_FILL
    ws[f"A{t}"].alignment = Alignment(horizontal="center")
    ws[f"F{t}"].number_format = "#,##0.00"
    ws[f"F{t}"].alignment = Alignment(horizontal="right")
    apply_main_layout(ws)

    d = wb.create_sheet("核对明细")
    d.append(["汇总序号", "乘坐人", "类型", "日期", "申报金额", "凭证金额", "支付金额", "凭证文件", "发票号码", "支付/乘车时间", "说明", "个人表来源"])
    _style_header(d, 1, 12)
    for r in records:
        ev = r["evidence"]
        it, pay, inv = ev.get("itinerary"), ev.get("payment"), ev.get("invoice")
        files = [x for x in [it and f"{it['file']}#第{it['no']}笔", pay and pay["file"]] + ev.get("supplements", []) if x]
        d.append([r["seq"], r["person"], r["type"], r["date"], r["amount"], it["amount"] if it else (pay["amount"] if pay else None), pay["amount"] if pay else None,
                  "\n".join(files), inv["number"] if inv else "", (pay["ride_at"] or pay["paid_at"]) if pay else (f"{it['date']} {it['time']}" if it else ""), "；".join(r.get("notes", [])), r.get("source", "")])
    for col, w in zip("ABCDEFGHIJKL", [8, 8, 7, 11, 9, 9, 9, 46, 22, 18, 30, 30]):
        d.column_dimensions[col].width = w
    for row in d.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    apply_page_setup(d, title_rows="1:1")

    q = wb.create_sheet("待确认事项")
    q.append(["乘坐人", "汇总序号", "级别", "问题", "处理建议"])
    _style_header(q, 1, 5)
    seq_of = {(r["person"], r["idx"]): r["seq"] for r in records}
    n = 0
    for p in sorted(persons, key=lambda p: pinyin_key(p["name"])):
        for x in p["report"]["issues"]:
            if x["level"] == "tip":
                continue
            q.append([p["name"], seq_of.get((p["name"], x["row"]), ""), "必须修改" if x["level"] == "error" else "需确认", x["message"], x["suggestion"]])
            n += 1
    if n == 0:
        q.append(["", "", "", "无待确认事项", ""])
    for col, w in zip("ABCDE", [8, 9, 9, 60, 50]):
        q.column_dimensions[col].width = w
    for row in q.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    apply_page_setup(q, title_rows="1:1")
    wb.save(path)
    return n


def build_month(month, persons, out_dir, cfg):
    """persons: [{name, base, report}]，report 中的文件路径相对 base。"""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    title = month_title(month)
    records = order_records(persons)
    xlsx = out_dir / f"{title}打车行程汇总.xlsx"
    pending = write_workbook(month, persons, records, xlsx, cfg["unit"])
    pack = out_dir / f"{title}_打印材料_已标序号.pdf"
    badges = build_print_pack(month, persons, records, pack)
    invpdf = out_dir / f"{title}_网约车电子发票合并.pdf"
    invoices = build_invoice_pack(persons, records, invpdf)
    taxipdf = out_dir / f"{title}_出租车支付记录汇总.pdf"
    build_taxi_pack(month, persons, records, taxipdf)
    stats = {
        "month": month,
        "generated_at": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "persons": [{"name": p["name"], "count": len(p["report"]["rows"]), "total": p["report"]["totals"]["all"], "passed": p["report"]["passed"]} for p in sorted(persons, key=lambda p: pinyin_key(p["name"]))],
        "count": len(records),
        "taxi": round(sum(r["amount"] or 0 for r in records if r["type"] == "出租车"), 2),
        "platform": round(sum(r["amount"] or 0 for r in records if r["type"] == "网约车"), 2),
        "total": round(sum(r["amount"] or 0 for r in records), 2),
        "pending": pending,
        "files": {"xlsx": xlsx.name, "print": pack.name, "invoices": invpdf.name, "taxi": taxipdf.name},
    }
    data = {"stats": stats, "records": records, "badges": badges, "invoices": invoices, "bases": {p["name"]: str(p["base"]) for p in persons}}
    (out_dir / "汇总数据.json").write_text(json.dumps(data, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return stats


def build_from_folders(input_dir, month, out_dir, cfg, cache_dir=None):
    """管理员批量模式：input_dir 下每个子文件夹是一个人的材料（文件夹名=姓名）。"""
    out_dir = Path(out_dir)
    rep_dir = out_dir / "检查报告"
    rep_dir.mkdir(parents=True, exist_ok=True)
    persons = []
    for sub in sorted(p for p in Path(input_dir).iterdir() if p.is_dir() and not p.name.startswith(".")):
        rep = check_submission(sub, sub.name, month, cfg, cache_dir=cache_dir, files=collect_files(sub))
        (rep_dir / f"{sub.name}.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        (rep_dir / f"{sub.name}.txt").write_text(report_to_text(rep), encoding="utf-8")
        persons.append({"name": sub.name, "base": str(sub), "report": rep})
    stats = build_month(month, persons, out_dir, cfg)
    stats["reports"] = {p["name"]: {"passed": p["report"]["passed"], **p["report"]["counts"]} for p in persons}
    return stats
