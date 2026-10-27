"""报销专用的模型 harness：看不懂的截图交给视觉模型、规则检查后的模型复核、按凭证自动修改汇总表。

原则：金额、日期、凭证对应关系只认程序从凭证里读到的数字；模型负责理解截图、地点和公务事由，
不能放过“金额不符、缺支付记录、重复报销”这类硬性错误。
"""
import hashlib
import json
import shutil
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .llm import LLMError, available, chat, image_part
from .rules import place_group
from .template import write_template

RULES = """你是北京中关村学院打车报销审核员。报销的是本人的公务打车费用，规则如下：
1. 网约车：必须有平台电子发票 PDF 和电子行程单 PDF；汇总表金额 = 行程单该笔金额；同一张发票里的每一笔行程都必须填进汇总表。
2. 出租车：每一笔必须有支付记录截图（支付记录是最终依据，纸质发票不检查）；金额 = 实际支付金额。
   用了优惠券：报优惠后的实付金额，并附显示优惠券抵扣的截图。含路桥费/高速费：不报销，需从支付金额中扣除并附车费明细截图。
3. 支付时间可以比乘车晚 1-3 天；网约车银行立减等小额优惠按发票金额报销。
4. 起止地方向必须与公务事由一致。例如 起点“海龙大厦”→终点“北京中关村学院”应写“从真机实验场地返回学院”，不能写“前往真机实验场地”。
5. 常见公务地点：学院（北京中关村学院、海淀大悦信息科技园 C 区各号楼）；真机实验场地（中关村海龙大厦）；合作单位（深势科技、北京理工大学、PlugandPlay 璞跃中国、智源大厦、北京大学等）；
   学院与地铁站（如永丰南）之间的接驳在事由能说明公务关联时可以接受。
6. 银行、派出所、出入境、医院、政务大厅、营业厅、4S 店等与公务无关的地点不能报销；住宅小区、商场、酒店、机场、火车站需要在事由里说明公务原因。
7. 事由要具体（去做什么），不能只写“打车”“出行”“办事”。"""

# 模型新增问题只有这几类可以拦住提交，其余降为“需要确认”，避免模型结论的随机性反复卡住同学
HARD_AI_CATEGORIES = {"forbidden_place", "duplicate", "amount"}
TABLE_EXT_ALL = {".xlsx", ".xlsm", ".csv", ".xls"}

# 模型可以判定为“可以接受”的规则问题；其余（金额、缺凭证、重复、抬头）只能由材料修改解决
SOFT_CODES = {"direction_mismatch", "place_suspicious", "reason_vague", "image_unknown", "payment_unused", "detail_unused", "platform_paid_diff", "row_person", "pdf_unknown", "file_ignored", "itinerary_parse"}


def _sha(obj):
    return hashlib.sha256(json.dumps(obj, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()[:24]


def _cached(cache_dir, key, fn):
    path = Path(cache_dir) / f"ai_{key}.json" if cache_dir else None
    if path and path.is_file():
        return json.loads(path.read_text(encoding="utf-8"))
    out = fn()
    if path:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(out, ensure_ascii=False), encoding="utf-8")
    return out


# ---------- 1. 视觉模型识别截图 ----------
VISION_PROMPT = """这是报销材料里的一张截图或照片，OCR 程序没能识别。请判断类型并提取信息，只输出 JSON：
{"kind": "payment"(支付记录/账单详情) | "detail"(车费明细/订单详情，含优惠券、路桥费等) | "taxi_invoice_photo"(出租车纸质发票照片) | "other",
 "amount": 实际支付金额（正数，没有则 null）, "paid_at": "YYYY-MM-DD HH:MM:SS" 或 null, "ride_at": 乘车时间或 null,
 "merchant": 平台或商户, "product": 商品说明, "orders": [订单号/交易单号字符串],
 "fare": 车费原价或 null, "coupon": 优惠券抵扣金额（正数）或 null, "toll": 路桥费/高速费或 null, "date": 行程日期 "YYYY-MM-DD" 或 null,
 "note": "一句话说明这张图是什么"}
年份不确定时按 {year} 年。看不清的字段填 null，不要猜。"""


def vision_hook(cfg, cache_dir, year):
    if not available(cfg):
        return None

    def hook(path, info):
        sha = hashlib.sha256(Path(path).read_bytes()).hexdigest()

        def call():
            data, _ = chat(cfg, [{"role": "user", "content": [{"type": "text", "text": VISION_PROMPT.replace("{year}", year)}, image_part(path)]}], model=cfg["llm"]["vision_model"], max_tokens=4000)
            return data

        try:
            v = _cached(cache_dir, f"vision_{sha}", call)
        except (LLMError, ValueError) as e:
            info["ai_error"] = str(e)
            return info
        out = dict(info)
        out["source"] = "vision"
        out["ai_note"] = v.get("note") or ""
        kind = v.get("kind")
        num = lambda k: round(float(v[k]), 2) if isinstance(v.get(k), (int, float)) and v[k] > 0 else None  # noqa: E731
        common = {"merchant": v.get("merchant") or info.get("merchant", ""), "product": v.get("product") or info.get("product", ""), "orders": [str(o) for o in (v.get("orders") or [])],
                  "paid_at": v.get("paid_at"), "ride_at": v.get("ride_at"), "ride_hint": None, "toll": num("toll"), "coupon": num("coupon"), "fare": num("fare")}
        if kind == "payment" and num("amount"):
            out.update(common, kind="payment", entries=[{"amount": num("amount"), "paid_at": v.get("paid_at"), "ride_at": v.get("ride_at")}], discount=num("coupon"))
        elif kind == "detail":
            out.update(common, kind="detail", pay_amount=num("amount"), date=v.get("date") or v.get("ride_at") or v.get("paid_at"))
        elif kind == "taxi_invoice_photo":
            out["kind"] = "taxi_invoice_photo"
        return out

    return hook


# ---------- 2. 模型复核 ----------
REVIEW_PROMPT = """下面是一位同学 {month} 的打车报销材料经程序解析和规则检查后的结果（JSON）。请你作为审核员复核：
- rows 是他汇总表里的每一行及程序匹配到的凭证；rule_issues 是规则程序发现的问题（带编号 id）；
  unrecognized_images 是程序看不懂的图片 OCR 文字；unused_payments 是没对上行程的支付记录。
- 只补充规则程序没发现的问题（例如事由与目的地明显无关、地点可疑但没被发现、同一行程重复填写、乘车时间不合常理、截图其实是别的东西等），不要重复已有问题。
- 每个问题给出 category：forbidden_place（非公务地点）、duplicate（同一行程/同一笔支付重复填写）、amount（金额与凭证明确不符）、other（其他，例如事由写法、地点名称不一致）。
  只有前三类且证据明确时才用 error；other 一律用 warning，由管理员判断。不要编造材料里没有的事实。
- rule_issues 中 soft=true 的问题，如果结合上下文你判断其实合理（例如地铁站接驳去实验场地），可以放进 dismiss 并说明理由；soft=false 的不能 dismiss。
只输出 JSON：{"summary": "一两句话的总体结论", "add_issues": [{"level": "error"|"warning", "category": "forbidden_place"|"duplicate"|"amount"|"other", "row": 行号或 null, "message": "问题", "suggestion": "具体怎么改"}], "dismiss": [{"id": 编号, "reason": "理由"}]}

材料：
{data}"""


def _review_input(rep):
    rows = []
    for r in rep["rows"]:
        ev = r["evidence"]
        e = {}
        if "itinerary" in ev:
            it = ev["itinerary"]
            e["行程单"] = f"{it['date']} {it['time']} {it['start']}→{it['end']} {it['amount']}元 {it['provider']}"
        if "payment" in ev:
            p = ev["payment"]
            e["支付记录"] = f"{p['file']} {p['amount']}元 支付{p['paid_at'] or '?'} 乘车{p['ride_at'] or '?'}"
        if ev.get("supplements"):
            e["补充截图"] = ev["supplements"]
        rows.append({"行": r["idx"], "日期": r["date"], "类型": r["type"], "起点": r["start"], "终点": r["end"], "事由": r["reason"], "金额": r["amount"], "备注": r.get("remark") or "", "凭证": e, "说明": r.get("notes", [])})
    issues = [{"id": k, "level": x["level"], "code": x["code"], "row": x["row"], "message": x["message"], "soft": x["code"] in SOFT_CODES} for k, x in enumerate(rep["issues"]) if x["level"] != "tip"]
    unknown = [{"file": im["file"], "ocr": im.get("text", "")[:1200], "ai_note": im.get("ai_note", "")} for im in rep["images"] if im["kind"] == "unknown"]
    used = {r["evidence"]["payment"]["file"] for r in rep["rows"] if "payment" in r["evidence"]}
    unused = [{"file": im["file"], "entries": im.get("entries"), "merchant": im.get("merchant")} for im in rep["images"] if im["kind"] == "payment" and im["file"] not in used]
    return {"name": rep["name"], "month": rep["month"], "totals": rep["totals"], "rows": rows, "rule_issues": issues, "unrecognized_images": unknown, "unused_payments": unused,
            "files": [f"{f['path']}（{f['kind']}）" for f in rep["files"]]}


def review(rep, cfg, cache_dir=None):
    """在报告上就地合并模型复核结果，返回 rep。"""
    if not available(cfg):
        rep["ai"] = {"status": "disabled"}
        return rep
    data = _review_input(rep)
    try:
        res = _cached(cache_dir, "review_" + _sha([data, RULES, REVIEW_PROMPT]), lambda: chat(cfg, [{"role": "system", "content": RULES}, {"role": "user", "content": REVIEW_PROMPT.replace("{month}", rep["month"]).replace("{data}", json.dumps(data, ensure_ascii=False))}])[0])
    except (LLMError, ValueError) as e:
        rep["ai"] = {"status": "error", "error": str(e)}
        rep["issues"].append({"level": "tip", "code": "ai_unavailable", "message": "模型复核暂时不可用，本次只做了规则检查", "suggestion": "", "row": None, "file": None})
        return finalize(rep)
    nrows = len(rep["rows"])
    dismissed = []
    for d in res.get("dismiss") or []:
        k = d.get("id")
        if isinstance(k, int) and 0 <= k < len(rep["issues"]) and rep["issues"][k]["code"] in SOFT_CODES and rep["issues"][k]["level"] != "tip":
            x = rep["issues"][k]
            x["level"] = "tip"
            x["ai_dismissed"] = d.get("reason", "")
            x["message"] += f"（模型复核认为可以接受：{d.get('reason', '')}）"
            dismissed.append(k)
    added = 0
    for a in (res.get("add_issues") or [])[:20]:
        if not a.get("message"):
            continue
        row = a.get("row") if isinstance(a.get("row"), int) and 1 <= a["row"] <= nrows else None
        hard = a.get("level") == "error" and a.get("category") in HARD_AI_CATEGORIES
        rep["issues"].append({"level": "error" if hard else "warning", "code": "ai_review", "message": "【模型复核】" + str(a["message"]), "suggestion": str(a.get("suggestion") or ""), "row": row, "file": None})
        added += 1
    rep["ai"] = {"status": "ok", "model": cfg["llm"]["model"], "summary": str(res.get("summary") or ""), "added": added, "dismissed": dismissed}
    return finalize(rep)


def finalize(rep):
    for r in rep["rows"]:
        r["issues"] = [k for k, x in enumerate(rep["issues"]) if x["row"] == r["idx"]]
    rep["counts"] = {lv: sum(1 for x in rep["issues"] if x["level"] == lv) for lv in ("error", "warning", "tip")}
    rep["passed"] = rep["counts"]["error"] == 0
    return rep


# ---------- 3. 按凭证修改汇总表 ----------
FIX_PROMPT = """下面是程序根据凭证整理出的修改版汇总表草稿（rows），以及原表和检查问题。请只修改 起点/终点/公务事由 三列，让每一行：
- 公务事由具体、与起止地方向一致（去程写“前往…”，返程写“从…返回学院”）；原事由正确就保持原样，只改有问题的；
- 事由为空的（多为补填的行程单行程），根据起止地推断一个合理的公务事由，并在 note 写“模型推断，请本人核对”；
- 网约车（route_locked=true）的起止地以行程单为准，不能改；出租车只有在起止地明显写反时才可以交换起点和终点；
- 不要改日期、金额、类型，不要增删行。
只输出 JSON：{"rows": [{"id": 行id, "start": "...", "end": "...", "reason": "...", "note": "改了什么（没改就空字符串）"}], "summary": "一句话"}

数据：
{data}"""


def _issues_by_row(rep):
    d = defaultdict(set)
    for x in rep["issues"]:
        if x["row"] and x["level"] != "tip":
            d[x["row"]].add(x["code"])
    return d


def _template_reason(start, end, cfg):
    eg, sg = place_group(end, cfg), place_group(start, cfg)
    if eg == "学院":
        return f"从{sg or start}返回学院"
    if eg == "实验场地":
        return "前往真机实验场地"
    return f"前往{end}" if end else ""


def draft_fix(rep, cfg):
    """按凭证生成修改草稿。返回 (rows, removed, notes)。rows 每行含 id/date/type/start/end/reason/amount/route_locked/changes。"""
    codes = _issues_by_row(rep)
    out, removed, notes = [], [], []
    used_it = set()
    for r in rep["rows"]:
        c = codes.get(r["idx"], set())
        ev = r["evidence"]
        row = {"date": r["date"], "type": r["type"], "start": r["start"], "end": r["end"], "reason": r["reason"] or "", "amount": r["amount"], "route_locked": False, "changes": [], "orig": r["idx"]}
        desc = f"原第{r['idx']}行（{r['date']} {r['start']}→{r['end']} {r['amount']}元）"
        if c & {"place_forbidden"}:
            removed.append(f"{desc}：地点与公务无关，已删除")
            continue
        if c & {"row_month", "row_date", "row_amount"}:
            removed.append(f"{desc}：日期/金额无法识别或不属于本月，已删除（如属其他月份请在对应月份提交）")
            continue
        if r["type"] == "网约车":
            if "itinerary" not in ev:
                removed.append(f"{desc}：在行程单里找不到，已删除（如确有该行程，请补充电子发票和行程单后再加回）")
                continue
            it = ev["itinerary"]
            used_it.add((it["file"], it["no"]))
            if (row["start"], row["end"]) != (it["start"], it["end"]):
                row["changes"].append("起止地按行程单填写")
            row.update(date=it["date"], start=it["start"], end=it["end"], amount=it["amount"], route_locked=True)
        else:
            if "payment" in ev:
                pass
            elif "near_payment" in ev and "taxi_more_than_paid" in c:
                p = ev["near_payment"]
                row["changes"].append(f"金额 {r['amount']} 改为实付 {p['amount']} 元（用了优惠券请附优惠截图）")
                row["amount"] = p["amount"]
            elif "taxi_less_than_paid" in c:
                notes.append(f"{desc}：实付比填写金额多，若差额是路桥费请补充车费明细截图，否则请按实付填写（未自动修改）")
            else:
                removed.append(f"{desc}：没有对应的支付记录，已删除（如确有支付，请补充支付截图后再加回）")
                continue
        if c & {"route_missing", "reason_missing", "reason_vague", "direction_reversed", "direction_mismatch", "direction_vs_itinerary", "place_suspicious", "ai_review"}:
            row["need_review"] = sorted(c)
        out.append(row)
    for it in rep["itineraries"]:
        if it.get("is_taxi"):
            continue
        for tr in it["rows"]:
            if (it["file"], tr["no"]) not in used_it:
                out.append({"date": tr["date"], "type": "网约车", "start": tr["start"], "end": tr["end"], "reason": "", "amount": tr["amount"], "route_locked": True,
                            "changes": [f"补填行程单 {Path(it['file']).name} 第{tr['no']}笔"], "orig": None, "need_review": ["reason_missing"]})
    out.sort(key=lambda x: (x["date"] or "", x["type"] != "出租车", x["orig"] or 999))
    for i, row in enumerate(out, 1):
        row["id"] = i
    for x in rep["issues"]:
        if x["code"] == "payment_unused":
            notes.append(f"{x['message']}——如果是公务打车，请在表中补一行并填写起止地和事由")
        elif x["code"] in ("ai_review", "duplicate_submission", "identical_submission") and x["level"] == "error":
            notes.append(f"{x['message']}——{x.get('suggestion') or '请本人确认后修改'}（未自动修改）")
    return out, removed, notes


def ai_fix(rep, rows, cfg, cache_dir=None):
    """让模型修改起止地/事由；校验后合并。返回 (rows, summary, used_model)。"""
    for row in rows:
        if not row["reason"]:
            row["reason"] = _template_reason(row["start"], row["end"], cfg)
            row["changes"].append("按起止地补写事由，请本人核对")
    if not available(cfg):
        return rows, "", False
    data = {"month": rep["month"], "rows": [{k: row.get(k) for k in ("id", "date", "type", "start", "end", "reason", "amount", "route_locked", "need_review")} for row in rows],
            "issues": [x["message"] for x in rep["issues"] if x["level"] != "tip"]}
    try:
        res = _cached(cache_dir, "fix_" + _sha([data, RULES, FIX_PROMPT]), lambda: chat(cfg, [{"role": "system", "content": RULES}, {"role": "user", "content": FIX_PROMPT.replace("{data}", json.dumps(data, ensure_ascii=False))}])[0])
    except (LLMError, ValueError):
        return rows, "", False
    byid = {row["id"]: row for row in rows}
    for m in res.get("rows") or []:
        row = byid.get(m.get("id"))
        if not row:
            continue
        start, end, reason = str(m.get("start") or row["start"]), str(m.get("end") or row["end"]), str(m.get("reason") or row["reason"]).strip()[:80]
        if row["route_locked"] or (start, end) not in ((row["start"], row["end"]), (row["end"], row["start"])):
            start, end = row["start"], row["end"]
        if (start, end) != (row["start"], row["end"]):
            row["changes"].append("交换起点和终点")
        if reason and reason != row["reason"]:
            row["changes"].append(f"事由“{row['reason'] or '空'}”改为“{reason}”")
        row.update(start=start, end=end, reason=reason)
        if m.get("note") and "核对" in str(m["note"]):
            row["changes"].append("请本人核对事由")
    return rows, str(res.get("summary") or ""), True


def write_table(rows, path, name):
    write_template(path)
    wb = load_workbook(path)
    ws = wb["打车明细"]
    for i, row in enumerate(rows, 2):
        vals = [i - 1, row["date"], row["type"], row["start"], row["end"], row["reason"], row["amount"], "；".join(row["changes"])]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v)
    wb.save(path)


def build_fix(rep, upload_dir, out_dir, cfg, cache_dir=None):
    """生成修改版材料目录和压缩包。返回 dict(dir, zip, rows, removed, notes, summary, used_model)。"""
    rows, removed, notes = draft_fix(rep, cfg)
    rows, summary, used_model = ai_fix(rep, rows, cfg, cache_dir)
    out_dir = Path(out_dir)
    mat = out_dir / "修改版材料"
    if mat.exists():
        shutil.rmtree(mat)
    mat.mkdir(parents=True)
    upload_dir = Path(upload_dir)
    for p in sorted(upload_dir.rglob("*")):
        if p.is_file() and p.suffix.lower() not in TABLE_EXT_ALL and not p.name.startswith("."):
            dst = mat / p.relative_to(upload_dir)
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(p, dst)
    table = mat / f"个人打车汇总表_{rep['name']}_{rep['month']}_修改版.xlsx"
    write_table(rows, table, rep["name"])
    lines = [f"{rep['name']} {rep['month']} 打车报销材料修改说明", "", "这份材料由系统按凭证自动修改，原汇总表已替换为“修改版”。请逐条核对，尤其是标注“请本人核对”的公务事由，确认属实后再提交。", ""]
    if summary:
        lines += [f"模型说明：{summary}", ""]
    changed = [r for r in rows if r["changes"]]
    lines.append(f"一、修改的行（{len(changed)} 行，见表中“备注”列）")
    lines += [f"  第{r['id']}行{'（原第' + str(r['orig']) + '行）' if r.get('orig') and r['orig'] != r['id'] else ''} {r['date']} {r['start']}→{r['end']} {r['amount']}元：{'；'.join(r['changes'])}" for r in changed] or ["  无"]
    lines += ["", f"二、删除的行（{len(removed)} 行）"] + ([f"  {x}" for x in removed] or ["  无"])
    lines += ["", "三、还需要你处理"] + ([f"  {x}" for x in notes] or ["  无"])
    lines += ["", f"修改后共 {len(rows)} 笔，合计 {round(sum(r['amount'] or 0 for r in rows), 2):.2f} 元。"]
    (mat / "修改说明.txt").write_text("\n".join(lines), encoding="utf-8")
    zpath = out_dir / f"{rep['name']}_{rep['month']}_修改版材料.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(mat.rglob("*")):
            if p.is_file():
                z.write(p, p.relative_to(mat))
    return {"dir": str(mat), "zip": str(zpath), "rows": rows, "removed": removed, "notes": notes, "summary": summary, "used_model": used_model, "text": "\n".join(lines)}
