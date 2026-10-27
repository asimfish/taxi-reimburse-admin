from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .layout import BORDER

HEADERS = ["序号", "日期", "类型", "起点", "终点", "公务事由", "金额", "备注"]
WIDTHS = [6, 12, 9, 28, 28, 30, 10, 24]
GUIDE = [
    "填写说明（本页不会被读取）",
    "1. 每一笔打车填一行；“类型”只填“出租车”或“网约车”。",
    "2. 日期写成 2026-07-14；金额写数字，出租车填写实际支付金额（用了优惠券填优惠后金额，路桥费不报销需扣除）。",
    "3. 网约车金额与电子行程单逐笔一致；同一张发票里的每一笔行程都要填写。",
    "4. 起点、终点写清楚地点；公务事由写去做什么，方向要和起止地一致，例如“从真机实验场地返回学院”对应 起点=海龙大厦、终点=北京中关村学院。",
    "5. 一起上传：本表、网约车电子发票 PDF + 电子行程单 PDF、出租车支付记录截图（每笔一张）。",
    "6. 支付金额与发票/行程单不同（优惠券、路桥费）时，再上传一张显示费用明细的截图，文件名写“原截图名+补充”。",
    "7. 建议：把出租车纸质发票按本表顺序贴在 A4 纸上拍照一起上传（纸质原件仍交给管理员）。",
]


def write_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "打车明细"
    ws.append(HEADERS)
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor="FF31465F")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
    dv = DataValidation(type="list", formula1='"出租车,网约车"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C2:C200")
    for r in range(2, 61):
        ws[f"B{r}"].number_format = "yyyy-mm-dd"
        ws[f"G{r}"].number_format = "0.00"
        for col in "ABCDEFGH":
            ws[f"{col}{r}"].border = BORDER
    ws.freeze_panes = "A2"
    g = wb.create_sheet("填写说明")
    g.column_dimensions["A"].width = 120
    for i, line in enumerate(GUIDE, 1):
        g[f"A{i}"] = line
        g[f"A{i}"].alignment = Alignment(wrap_text=True)
    g["A1"].font = Font(bold=True, size=12)
    wb.save(path)
